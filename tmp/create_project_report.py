from pathlib import Path
from textwrap import wrap

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


ROOT = Path('/Users/sreerajmuthaiya.a.l/llvm-vectorization')
SOURCE = Path('/Users/sreerajmuthaiya.a.l/Downloads/Project - Report Template.docx')
OUTPUT = ROOT / 'output/Rust_Loop_Vectorizer_Project_Report.docx'
FIGURE = ROOT / 'tmp/report-cli-figure.png'


def set_text(paragraph, value):
    if paragraph.runs:
        paragraph.runs[0].text = value
        for run in paragraph.runs[1:]:
            run.text = ''
    else:
        paragraph.add_run(value)


def set_font(paragraph, size=12, bold=False):
    for run in paragraph.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(size)
        run.bold = bold


def insert_body(before, text, *, bold_lead=None):
    p = before.insert_paragraph_before(style='Normal')
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.left_indent = Inches(0.24)
    p.paragraph_format.right_indent = Inches(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.15
    if bold_lead:
        lead = p.add_run(bold_lead)
        lead.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    set_font(p)
    return p


def remove_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)


def create_figure():
    line = (ROOT / 'tmp/report-cli.txt').read_text().splitlines()[0]
    fields = dict(part.split('=', 1) for part in line.split() if '=' in part)
    ir = (ROOT / 'tmp/report-vectorized.ll').read_text().splitlines()
    wide = []
    for needle in ('load <4 x float>', 'fadd <4 x float>', 'store <4 x float>'):
        wide.append(next(x.strip() for x in ir if needle in x))
    lines = [
        '$ target/release/rv-vectorize --report tests/fixtures/vectorizable.ll -o output.ll',
        f"function={fields['function']}  decision={fields['decision']}  vf={fields['vf']}",
        f"estimated_vector_coverage={fields['estimated_vector_coverage']}  issued_lane_utilization={fields['issued_lane_utilization']}",
        '',
        '$ selected generated LLVM IR',
    ]
    lines += wide
    font_path = '/System/Library/Fonts/Menlo.ttc'
    font = ImageFont.truetype(font_path, 23)
    image = Image.new('RGB', (1500, 350), '#F4F4F4')
    draw = ImageDraw.Draw(image)
    draw.rounded_rectangle((4, 4, 1495, 345), radius=12, outline='#777777', width=2, fill='#F7F7F7')
    y = 22
    for i, text in enumerate(lines):
        if text:
            draw.text((30, y), text, fill='#111111', font=font)
        y += 39 if i != 3 else 28
    image.save(FIGURE)


doc = Document(SOURCE)
set_text(doc.paragraphs[3], 'Rust Loop Vectorizer for LLVM 21')
set_text(doc.paragraphs[11], 'October 2026')

registrations = ['24BCE0250', '24BCE0016', '24BAI0340']
names = ['Sree Raj Muthaiya.A.L', 'Karnati Abhinav Reddy', 'Nidhurshek Prabakaran']
for ci, values in enumerate([registrations, names]):
    cell = doc.tables[0].cell(0, ci)
    for pi, value in enumerate(values, 1):
        if pi < len(cell.paragraphs):
            p = cell.paragraphs[pi]
            set_text(p, value)
        else:
            p = cell.add_paragraph(value)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(1)
        p.paragraph_format.space_after = Pt(4)
        set_font(p, 10.5, True)
    for p in cell.paragraphs[1 + len(values):]:
        set_text(p, '')
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

headings = {
    p.text.strip(): p for p in doc.paragraphs
    if p.text.strip() in {
        'Abstract:', 'Introduction:', 'Novelty in Project:', 'Proposed Work:',
        'Implementation & Screenshots:', 'Conclusion:',
        'Project Repository Link: (GitHub or drive link)'
    }
}
last = headings['Project Repository Link: (GitHub or drive link)']
set_text(last, 'Project Repository Link:')

# The empty paragraphs are demonstration space in the reference, not report text.
after_abstract = False
for p in list(doc.paragraphs):
    if p._element is headings['Abstract:']._element:
        after_abstract = True
    elif after_abstract and p.text == '':
        # Keep cover spacing and the blank page transition, but remove unused
        # body spacers after the first report heading.
        remove_paragraph(p)

insert_body(headings['Introduction:'],
    'This report presents a Rust implementation of a bounded automatic loop vectorizer for LLVM 21. It explains the supported loops, legality and cost decisions, LLVM IR generation, validation, and recorded benchmark result. The pass emits fixed-width SIMD operations only when its checks establish safety; otherwise it keeps the scalar loop. In one recorded Apple M5 benchmark, the custom variant achieved a 3.009× geometric-mean speedup over scalar execution across seven kernels, with the result limited to that host and corpus.')

insert_body(headings['Novelty in Project:'],
    'This project implements an automatic loop vectorization pass for LLVM 21 in Rust. It identifies a deliberately bounded class of loops, proves that their memory operations can be reordered safely, chooses a fixed vector width, and emits vector LLVM IR while retaining the scalar loop for short inputs and any remainder. The command-line driver accepts LLVM IR, bitcode, supported source files, or a project directory.')
insert_body(headings['Novelty in Project:'],
    'Loop vectorization combines several scalar iterations into one SIMD operation. A correct implementation must preserve data dependences before it asks whether vector execution is profitable. The project therefore separates discovery and legality checks from width selection and IR rewriting. Unsupported or uncertain loops remain scalar and receive a reason in report mode.')

insert_body(headings['Proposed Work:'],
    'The main contribution is a correctness-first LLVM pass whose discovery, dependence classification, cost planning, and IR construction live in Rust. A small C++ adapter registers the pass with LLVM’s New Pass Manager and exposes only the LLVM operations required by the Rust implementation. The pass does not rely on an unchecked source-level pattern match: it validates the loop structure, data flow, memory layout, and alias relationships before modifying IR.')
insert_body(headings['Proposed Work:'],
    'The tool also makes decisions observable. It reports the chosen vector factor, estimated vector coverage, issued-lane utilization, rejection reason, and per-loop analysis and transformation time. Its conservative, balanced, and aggressive policies expose the trade-off between accepting a legal loop and predicting a useful speedup. A terminal interface and project-directory mode make these decisions easier to inspect across files.')

insert_body(headings['Implementation & Screenshots:'],
    'The implemented workflow follows five stages. First, the driver obtains LLVM 21 IR from a direct input or a supported frontend. Second, the pass finds canonical one-block loops with a zero-based, unit-step i64 induction variable. Third, it checks supported operations, affine memory accesses, target data layout, aliasing, and loop-carried dependences. Fourth, a transparent cost model selects a fixed vector factor when the estimated trip count and gain meet the selected policy. Fifth, the transformer creates a vector loop and connects the original scalar loop as fallback and tail; LLVM verification checks the resulting module.')
insert_body(headings['Implementation & Screenshots:'],
    'The accepted subset includes contiguous loads and stores and lane-wise arithmetic, casts, comparisons, and select operations. Calls, reductions, multi-block loops, unresolved pointer aliases, and nonzero loop-carried memory distances are rejected by the Rust pass. These boundaries keep its safety claim precise. Future work can add stronger translation validation, LLVM analysis integration, runtime alias checks, reductions, and broader target evaluation.')

insert_body(headings['Conclusion:'],
    'The code is organized around src/vectorizer.rs for loop analysis and rewriting, src/dependence.rs for affine dependence classification, src/cost.rs for vector-factor planning, and native/pass_plugin.cpp for LLVM registration. The rv-vectorize driver runs the pass, LLVM verifier, and optional reporting; rv-vectorize-tui provides an interactive terminal view. The output IR contains real vector loads, arithmetic, and stores, with a scalar remainder path.')
insert_body(headings['Conclusion:'],
    'Figure 1 records a local run of the checked-in release binary on the repository’s positive LLVM fixture. The report marks add_f32 vectorized at VF 4, and the output IR contains <4 x float> load, fadd, and store instructions. The displayed coverage and lane-use values are compiler estimates for that loop; they are not measured hardware occupancy.')

create_figure()
fig = headings['Conclusion:'].insert_paragraph_before(style='Normal')
fig.alignment = WD_ALIGN_PARAGRAPH.CENTER
fig.paragraph_format.space_before = Pt(4)
fig.paragraph_format.space_after = Pt(2)
fig.add_run().add_picture(str(FIGURE), width=Inches(5.9))
cap = headings['Conclusion:'].insert_paragraph_before(style='Normal')
cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
cap.paragraph_format.space_after = Pt(8)
cap.add_run('Figure 1. CLI decision and generated vector IR from a local fixture run')
set_font(cap, 10)

insert_body(last,
    'Validation includes Rust tests, accepted and rejected LLVM fixtures, IR verification, native differential execution, and guard-page tail checks. A recorded 8 September 2026 run on one Apple M5 with LLVM 21.1.3 passed correctness checks for all seven scalar, custom, and LLVM kernels. The geometric-mean speedups over scalar execution were 3.009× for this pass and 2.974× for LLVM LoopVectorize. On 128 cloned loops, this pass’s in-process analysis plus transformation measured 2.355 µs at p50, 8.790 µs at p95, and 31.404 µs at p99. These results apply to the recorded corpus and host.')
insert_body(last,
    'The pass shows that a bounded legality contract can produce vector code with clear rejection reasons. Its current scope is canonical affine loops; stronger semantic validation should precede wider loop support.')

repo = doc.add_paragraph(style='Normal')
repo.alignment = WD_ALIGN_PARAGRAPH.LEFT
repo.paragraph_format.left_indent = Inches(0.24)
repo.paragraph_format.space_before = Pt(3)
repo.add_run('https://github.com/Passive-Coder/llvm-vectorization')
set_font(repo, 11)

# Move repository URL directly under its numbered heading if the reference had
# trailing blank paragraphs or section properties.
last._element.addnext(repo._element)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
