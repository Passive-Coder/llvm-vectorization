from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path("/Users/sreerajmuthaiya.a.l/llvm-vectorization/output/Rust_Loop_Vectorizer_Five_Speaker_Script.docx")

SECTIONS = [
    (
        "Speaker 1 Purpose",
        [
            "Good morning everyone. Our project is the Rust Loop Vectorizer for LLVM 21. Its purpose is to speed up suitable loops by processing several data elements with vector instructions. Imagine adding two arrays. A scalar loop adds one pair of numbers at a time; a vectorized loop can add several pairs together. This is useful when the same calculation repeats across many elements.",
            "We built the project as an LLVM optimization pass. It works on LLVM’s intermediate representation before the compiler produces machine code. The project is also a research and learning tool: it shows how a compiler can find parallel work, decide whether an optimization is safe, and generate faster code without changing the program’s result.",
        ],
    ),
    (
        "Speaker 2 How it works",
        [
            "The implementation has two main parts. A small C++ bridge registers the pass with LLVM, while Rust handles loop discovery, safety checks, planning, and the generation of new LLVM instructions.",
            "For each candidate loop, the pass checks its structure and memory accesses. It then estimates whether vectorization is worthwhile and chooses a fixed vector width. When target information is available, that choice can use the target’s vector register width.",
            "If the loop is accepted, the pass creates a vector loop for full groups of elements. The original scalar loop remains available for short inputs and for any elements left over at the end. This ensures that every element is processed exactly once.",
        ],
    ),
    (
        "Speaker 3 Features and safety",
        [
            "Correctness is central to the design. The Rust pass currently focuses on simple, single-block, count-up loops with contiguous memory access. It checks whether memory locations might overlap and whether one loop iteration depends on another. If it cannot establish safety, it leaves the loop unchanged and reports the reason.",
            "For accepted loops, it can generate vector loads and stores, arithmetic, casts, comparisons, and select operations. It supports common integer and floating-point element types. It also checks the target’s memory layout before changing memory instructions.",
            "More complex cases, including reductions and loops with multiple blocks, are outside the Rust pass’s current scope. That boundary matters: a performance improvement is only useful if the program still computes the correct result.",
        ],
    ),
    (
        "Speaker 4 Tools for using the project",
        [
            "We also built tools around the pass. The command-line application accepts LLVM IR, bitcode, supported source files, or a project directory. It can produce transformed IR, bitcode, or object files and report what happened to each candidate loop.",
            "Users can select conservative, balanced, or aggressive policies. They can also request a fixed vector width while retaining the safety checks. A terminal interface lets users change settings, run the pass, and inspect a side-by-side difference.",
            "For larger projects, the application can use a C-family compilation database or a manifest of existing LLVM IR files. It also supports JSON reports, parallel processing, and caching. A hybrid mode can run LLVM’s own Loop Vectorizer after the Rust pass for broader loop coverage.",
        ],
    ),
    (
        "Speaker 5 Testing and conclusion",
        [
            "Finally, we tested both correctness and performance. The project includes unit tests, LLVM verification, tests for loops that must be rejected, and native comparisons between scalar and vectorized results.",
            "In the recorded seven-kernel benchmark on an Apple M5, all variants passed the correctness checks. The custom pass achieved a geometric-mean speedup of about 3.01 times over the scalar version. LLVM’s built-in Loop Vectorizer achieved about 2.97 times on the same kernels. In a separate set of 128 cloned loops, the recorded median time for the custom pass’s in-process analysis and transformation was 2.355 microseconds per loop.",
            "These are measurements from one test setup, not a guarantee for every program. Overall, our project demonstrates a working vectorization pipeline with measurable results, clear safety limits, and room for future extensions. Thank you.",
        ],
    ),
]


def main() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(0.78)
    sec.bottom_margin = Inches(0.76)
    sec.left_margin = Inches(0.9)
    sec.right_margin = Inches(0.9)

    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.8)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.line_spacing = 1.18
    normal.paragraph_format.space_after = Pt(7)

    title_style = doc.styles["Title"]
    title_style.font.name = "Arial"
    title_style.font.size = Pt(22)
    title_style.font.bold = True
    title_style.font.color.rgb = RGBColor(0, 0, 0)
    title_style.paragraph_format.space_after = Pt(5)
    title_border = title_style._element.pPr.find(qn("w:pBdr"))
    if title_border is not None:
        title_style._element.pPr.remove(title_border)

    subtitle_style = doc.styles["Subtitle"]
    subtitle_style.font.name = "Arial"
    subtitle_style.font.size = Pt(10.5)
    subtitle_style.font.italic = False
    subtitle_style.font.color.rgb = RGBColor(0, 0, 0)
    subtitle_style.paragraph_format.space_after = Pt(17)

    heading_style = doc.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style.font.size = Pt(12)
    heading_style.font.bold = True
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.paragraph_format.space_before = Pt(12)
    heading_style.paragraph_format.space_after = Pt(5)
    heading_style.paragraph_format.keep_with_next = True

    title = doc.add_paragraph(style="Title")
    title.add_run("Rust Loop Vectorizer for LLVM 21")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("Five speaker presentation script for a five minute group talk")

    for heading, paragraphs in SECTIONS:
        doc.add_paragraph(heading, style="Heading 1")
        for index, text in enumerate(paragraphs):
            paragraph = doc.add_paragraph(text, style="Normal")
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = index < len(paragraphs) - 1

    doc.core_properties.title = "Rust Loop Vectorizer for LLVM 21"
    doc.core_properties.subject = "Five speaker presentation script"
    doc.core_properties.keywords = "LLVM, Rust, loop vectorization, presentation"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
